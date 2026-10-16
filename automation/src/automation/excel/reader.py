"""Reads test cases from the workbook through the project's column map.

Columns are found by header text from ``project.yaml``, never by position, so
QA can reorder or add columns freely. A missing sheet or required column is
fatal (``WorkbookError``); a bad value in one row becomes a ``ValidationIssue``
for that case only.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from automation.core.config import WorkbookSettings
from automation.core.exceptions import WorkbookError
from automation.excel.models import (
    AutomationRef,
    AutomationRefError,
    Severity,
    TestCase,
    ValidationIssue,
)

REQUIRED_CASE_FIELDS = (
    "id",
    "module",
    "feature",
    "scenario",
    "priority",
    "test_type",
    "preconditions",
    "steps",
    "expected_result",
    "automation_status",
)
REQUIRED_MAPPING_FIELDS = ("id", "automation_id")
_LIST_SPLIT = re.compile(r"[,;\n]")


@dataclass
class WorkbookData:
    cases: list[TestCase]
    allowed_values: dict[str, list[str]]
    issues: list[ValidationIssue] = field(default_factory=list)


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def cell(row: tuple[Any, ...], index: int) -> str:
    """Text of one cell; read-only rows can be shorter than the header row."""
    return text(row[index]) if index < len(row) else ""


def split_list(value: str) -> tuple[str, ...]:
    return tuple(part.strip() for part in _LIST_SPLIT.split(value) if part.strip())


class WorkbookReader:
    def __init__(self, path: Path, settings: WorkbookSettings) -> None:
        self._path = path
        self._settings = settings

    def read(self) -> WorkbookData:
        if not self._path.is_file():
            raise WorkbookError(f"Workbook not found: {self._path}")
        workbook = load_workbook(self._path, read_only=True, data_only=True)
        try:
            issues: list[ValidationIssue] = []
            refs, levels = self._read_mapping(workbook, issues)
            cases = self._read_cases(workbook, refs, levels)
            return WorkbookData(cases, self._read_allowed_values(workbook), issues)
        finally:
            workbook.close()

    def _sheet(self, workbook: Any, role: str) -> Worksheet | None:
        name = self._settings.sheets.get(role)
        if name is None:
            return None
        if name not in workbook.sheetnames:
            raise WorkbookError(f"Sheet {name!r} ({role}) is missing from {self._path.name}.")
        return workbook[name]

    def _columns(self, sheet: Worksheet, role: str, required: tuple[str, ...]) -> dict[str, int]:
        """Field → 0-based column index, from the header row."""
        headers = {text(c.value): i for i, c in enumerate(next(sheet.iter_rows(max_row=1)))}
        mapping = self._settings.columns.get(role, {})
        missing_config = [f for f in required if f not in mapping]
        if missing_config:
            raise WorkbookError(f"project.yaml has no column for {role}.{missing_config}.")
        located = {f: headers[h] for f, h in mapping.items() if h in headers}
        missing = [mapping[f] for f in required if f not in located]
        if missing:
            raise WorkbookError(f"Sheet {sheet.title!r} is missing required column(s) {missing}.")
        return located

    def _rows(self, sheet: Worksheet):
        for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if any(text(v) for v in row):
                yield number, row

    def _read_mapping(
        self, workbook: Any, issues: list[ValidationIssue]
    ) -> tuple[dict[str, tuple[AutomationRef, ...]], dict[str, str]]:
        sheet = self._sheet(workbook, "automation_mapping")
        if sheet is None:
            return {}, {}
        cols = self._columns(sheet, "automation_mapping", REQUIRED_MAPPING_FIELDS)
        refs: dict[str, tuple[AutomationRef, ...]] = {}
        levels: dict[str, str] = {}
        for number, row in self._rows(sheet):
            case_id = cell(row, cols["id"])
            parsed = []
            for raw in split_list(cell(row, cols["automation_id"])):
                try:
                    parsed.append(AutomationRef.parse(raw))
                except AutomationRefError as error:
                    issues.append(
                        ValidationIssue(
                            Severity.ERROR, str(error), case_id, number, "Automation ID"
                        )
                    )
            refs[case_id] = tuple(parsed)
            if "automation_level" in cols:
                levels[case_id] = cell(row, cols["automation_level"])
        return refs, levels

    def _read_cases(
        self,
        workbook: Any,
        refs: dict[str, tuple[AutomationRef, ...]],
        levels: dict[str, str],
    ) -> list[TestCase]:
        sheet = self._sheet(workbook, "test_cases")
        if sheet is None:
            raise WorkbookError("project.yaml must name the test_cases sheet.")
        cols = self._columns(sheet, "test_cases", REQUIRED_CASE_FIELDS)
        cases = []
        for number, row in self._rows(sheet):
            value = {f: cell(row, i) for f, i in cols.items()}
            case_id = value["id"]
            cases.append(
                TestCase(
                    id=case_id,
                    row=number,
                    **{f: value[f] for f in REQUIRED_CASE_FIELDS if f != "id"},
                    test_data_ids=split_list(value.get("test_data_ids", "")),
                    suites=tuple(s.lower() for s in split_list(value.get("suites", ""))),
                    environments=tuple(
                        e.lower() for e in split_list(value.get("environments", ""))
                    ),
                    automation_level=levels.get(case_id, ""),
                    automation_refs=refs.get(case_id, ()),
                    api_endpoint=value.get("api_endpoint", ""),
                    triage_hint=value.get("triage_hint", ""),
                    extra={f: v for f, v in value.items() if f not in REQUIRED_CASE_FIELDS},
                )
            )
        return cases

    def _read_allowed_values(self, workbook: Any) -> dict[str, list[str]]:
        sheet = self._sheet(workbook, "dropdowns")
        if sheet is None or not self._settings.allowed_values:
            return {}
        columns: dict[str, list[str]] = {}
        # Read-only sheets cannot iterate columns; transpose the rows instead.
        rows = list(sheet.iter_rows(values_only=True))
        width = max((len(r) for r in rows), default=0)
        for index in range(width):
            header, *values = (row[index] if index < len(row) else None for row in rows)
            columns[text(header)] = [text(v) for v in values if text(v)]
        allowed = {}
        for field_name, header in self._settings.allowed_values.items():
            if header not in columns:
                raise WorkbookError(f"Dropdowns sheet has no list {header!r} ({field_name}).")
            allowed[field_name] = columns[header]
        return allowed
