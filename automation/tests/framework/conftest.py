"""A throwaway project profile and workbook per test, so self-tests never depend
on a real app."""

from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

PROJECT_YAML = """\
name: demo
display_name: Demo App
app_root: ../..
default_environment: uat
workbook:
  path: qa/cases.xlsx
  sheets:
    test_cases: Cases
    automation_mapping: Mapping
    dropdowns: Lists
  columns:
    test_cases:
      id: ID
      module: Module
      feature: Feature
      scenario: Title
      priority: Priority
      test_type: Type
      preconditions: Pre
      test_data_ids: Data
      steps: Steps
      expected_result: Expected
      automation_status: Status
      suites: Suites
      environments: Envs
    automation_mapping:
      id: ID
      automation_id: Automation ID
  allowed_values:
    priority: Priority
    test_type: Type
    automation_status: Status
    suites: Suite
  id_pattern: "^TC-[A-Z]+-[0-9]+$"
secrets:
  DEMO_API_BASE_URL: gateway host
  DEMO_PASSWORD: test account password
masking:
  extra_keys: [employeeCode]
"""

ENV_YAML = """\
name: {name}
protected: {protected}
api:
  base_url_env: DEMO_API_BASE_URL
  api_prefix: /api/v1
"""

CASE_HEADERS = ["ID", "Module", "Feature", "Title", "Priority", "Type", "Pre", "Data",
                "Steps", "Expected", "Status", "Suites", "Envs"]  # fmt: skip
LISTS = {
    "Priority": ["P1", "P2", "P3"],
    "Type": ["Functional", "Negative"],
    "Status": ["Automated", "Partial", "Planned", "Manual only"],
    "Suite": ["smoke", "regression", "uat"],
}


def case_row(case_id: str = "TC-AUTH-001", **overrides: Any) -> dict[str, Any]:
    row = {
        "ID": case_id,
        "Module": "Authentication",
        "Feature": "Login",
        "Title": f"Scenario {case_id}",
        "Priority": "P1",
        "Type": "Functional",
        "Pre": "Signed out",
        "Data": "TD-001, TD-002",
        "Steps": "1. Sign in",
        "Expected": "Lands on home",
        "Status": "Automated",
        "Suites": "smoke, regression",
        "Envs": "",
    }
    row.update(overrides)
    return row


def write_workbook(
    path: Path,
    cases: list[dict[str, Any]],
    mapping: dict[str, str],
    headers: list[str] = CASE_HEADERS,
    lists: dict[str, list[str]] = LISTS,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cases"
    sheet.append(headers)
    for case in cases:
        sheet.append([case.get(h) for h in headers])
    mapping_sheet = workbook.create_sheet("Mapping")
    mapping_sheet.append(["ID", "Automation ID"])
    for case_id, ref in mapping.items():
        mapping_sheet.append([case_id, ref])
    list_sheet = workbook.create_sheet("Lists")
    for column, (header, values) in enumerate(lists.items(), start=1):
        for row, value in enumerate([header, *values], start=1):
            list_sheet.cell(row=row, column=column, value=value)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


@pytest.fixture(autouse=True)
def isolated_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Self-tests never write into the real reports/ or history/ folders."""
    monkeypatch.setenv("AUTOMATION_REPORTS_DIR", str(tmp_path / "reports"))
    monkeypatch.setenv("AUTOMATION_HISTORY_DIR", str(tmp_path / "history"))


@pytest.fixture
def projects_root(tmp_path: Path) -> Path:
    root = tmp_path / "projects"
    project = root / "demo"
    (project / "environments").mkdir(parents=True)
    (project / "project.yaml").write_text(PROJECT_YAML, encoding="utf-8")
    for name, protected in (("uat", "false"), ("production", "true")):
        (project / "environments" / f"{name}.yaml").write_text(
            ENV_YAML.format(name=name, protected=protected), encoding="utf-8"
        )
    return root


@pytest.fixture
def project_dir(projects_root: Path) -> Path:
    return projects_root / "demo"


@pytest.fixture
def app_root(projects_root: Path) -> Path:
    """``app_root: ../..`` from the profile resolves to the tmp folder."""
    return projects_root.parent


@pytest.fixture
def make_workbook(app_root: Path) -> Callable[..., Path]:
    def make(cases: list[dict[str, Any]], mapping: dict[str, str], **kwargs: Any) -> Path:
        return write_workbook(app_root / "qa" / "cases.xlsx", cases, mapping, **kwargs)

    return make


@pytest.fixture
def dart_test(app_root: Path) -> Callable[[str], str]:
    """Creates an (empty) Dart test file so file-existence checks pass."""

    def make(relative: str) -> str:
        path = app_root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.touch()
        return relative

    return make
